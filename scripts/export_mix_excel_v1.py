from __future__ import annotations

import argparse
import hashlib
import json
import shutil
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPORTER_VERSION = "export_mix_excel_v1.py@0.1"
EXPECTED_HEADERS = ("视频制作标题", "视频制作口播内容")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def copy_cell_style(source: Any, target: Any) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy(source.alignment)
    if source.protection:
        target.protection = copy(source.protection)


def export_mix_excel(
    batch_path: Path,
    template_path: Path,
    output_path: Path,
    fixture_only: bool = False,
) -> dict[str, Any]:
    batch_path = batch_path.expanduser().resolve()
    template_path = template_path.expanduser().resolve()
    output_path = output_path.expanduser().resolve()
    for path in (batch_path, template_path):
        if not path.is_file():
            raise FileNotFoundError(path)
    if output_path == template_path:
        raise RuntimeError("Excel output must not overwrite the source Template.")
    batch = json.loads(batch_path.read_text(encoding="utf-8"))
    if batch.get("status") != "approved":
        raise RuntimeError("Excel Export requires an Approved Generation Batch.")
    if batch.get("profile") != "mix":
        raise RuntimeError("Mix Excel Export rejects non-mix batches.")
    if batch.get("human_review", {}).get("all_export_items_human_approved") is not True:
        raise RuntimeError("Every exported item must be Human Approved.")
    contents = batch.get("contents") or []
    if not contents or any(item.get("status") != "human_approved" for item in contents):
        raise RuntimeError("Approved Batch contains a non-approved export item.")
    if output_path.exists():
        raise RuntimeError("Excel output already exists and cannot be overwritten.")

    template_sha_before = sha256_file(template_path)
    workbook = load_workbook(template_path)
    sheet = workbook.active
    headers = (sheet.cell(4, 1).value, sheet.cell(4, 2).value)
    if headers != EXPECTED_HEADERS:
        workbook.close()
        raise RuntimeError(
            f"Mix Excel Row 4 headers mismatch: expected {EXPECTED_HEADERS!r}, got {headers!r}."
        )
    workbook.close()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, output_path)
    workbook = load_workbook(output_path)
    sheet = workbook.active
    style_title = copy(sheet.cell(5, 1)._style)
    style_narration = copy(sheet.cell(5, 2)._style)
    for row in range(5, max(5, sheet.max_row) + 1):
        for column in range(1, max(2, sheet.max_column) + 1):
            sheet.cell(row, column).value = None
    for index, item in enumerate(contents, start=5):
        title = str(item.get("title") or "").strip()
        narration = str(item.get("narration") or "").strip()
        if not title or not narration:
            workbook.close()
            raise RuntimeError("Excel title and narration are required.")
        title_cell = sheet.cell(index, 1, title)
        narration_cell = sheet.cell(index, 2, narration)
        title_cell._style = copy(style_title)
        narration_cell._style = copy(style_narration)
    workbook.save(output_path)
    workbook.close()

    validation_book = load_workbook(output_path, read_only=True, data_only=False)
    validation_sheet = validation_book.active
    if (
        validation_sheet.cell(4, 1).value,
        validation_sheet.cell(4, 2).value,
    ) != EXPECTED_HEADERS:
        validation_book.close()
        raise RuntimeError("Exported Mix Excel headers changed.")
    for index, item in enumerate(contents, start=5):
        if validation_sheet.cell(index, 1).value != item["title"]:
            validation_book.close()
            raise RuntimeError("Exported Mix Excel title mismatch.")
        if validation_sheet.cell(index, 2).value != item["narration"]:
            validation_book.close()
            raise RuntimeError("Exported Mix Excel narration mismatch.")
        for column in range(3, validation_sheet.max_column + 1):
            if validation_sheet.cell(index, column).value is not None:
                validation_book.close()
                raise RuntimeError("Internal metadata leaked into Mix Excel.")
    validation_book.close()
    template_sha_after = sha256_file(template_path)
    if template_sha_after != template_sha_before:
        raise RuntimeError("Source Excel Template was modified.")
    return {
        "schema_version": "mix-excel-export-receipt-v1.0",
        "exporter_version": EXPORTER_VERSION,
        "fixture_only": fixture_only,
        "batch_path": str(batch_path),
        "batch_sha256": sha256_file(batch_path),
        "template_path": str(template_path),
        "template_sha256": template_sha_before,
        "output_path": str(output_path),
        "output_sha256": sha256_file(output_path),
        "exported_row_count": len(contents),
        "row_4_headers": list(EXPECTED_HEADERS),
        "first_data_row": 5,
        "internal_metadata_columns_added": False,
        "source_template_modified": False,
        "validation_passed": True,
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Export a Human Approved Mix Script Batch into the frozen Excel Contract."
    )
    parser.add_argument("--batch", required=True)
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--fixture-only", action="store_true")
    args = parser.parse_args()
    receipt = export_mix_excel(
        Path(args.batch),
        Path(args.template),
        Path(args.output),
        args.fixture_only,
    )
    receipt_path = Path(args.output).expanduser().resolve().with_suffix(
        ".export_receipt.json"
    )
    if receipt_path.exists():
        raise RuntimeError("Excel export receipt already exists.")
    receipt_path.write_text(
        json.dumps(receipt, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("MIX EXCEL EXPORT V1 PASS")
    print(f"Rows: {receipt['exported_row_count']}")
    print(f"Fixture only: {receipt['fixture_only']}")
    print(f"Output: {receipt['output_path']}")
    print(f"Receipt: {receipt_path}")


if __name__ == "__main__":
    main()
