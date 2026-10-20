from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import os
import zipfile
import uuid
from copy import copy
from pathlib import Path
from typing import Any

EXPORTER_VERSION = "export_mix_excel_v1.py@0.1"
EXPECTED_HEADERS = ("视频制作标题", "视频制作口播内容")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def copy_cell_style(source: Any, target: Any) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy(source.alignment)
    if source.protection:
        target.protection = copy(source.protection)


from profile_compatibility_consumers_v1_1 import guarded as compatibility_guarded
import delivery_receipt_recovery_v1 as partial
import generation_recovery_v1 as recovery


@compatibility_guarded
def prepare_mix_export(*, pipeline_root, request_id, batch_path, template_path, output_path, authorization_context):
    root=Path(pipeline_root).resolve();batch_path=Path(batch_path).resolve()
    template_path=Path(template_path).resolve();output_path=Path(output_path).resolve()
    directory=recovery.request_root(root,request_id)
    intent_path,_=partial.export_paths(root,request_id)
    intent=partial.check(partial.read(intent_path),partial.EXPORT) if intent_path.is_file() else None
    proof=partial.identity(authorization_context,intent)
    if output_path.parent!=root/'output' or template_path!=root/'output/素材混剪&数字人口播混剪文案导入模板.xlsx':
        raise partial.ReceiptRecoveryError('DELIVERY_SOURCE_CHANGED','Output and frozen template must belong to this runtime')
    if not intent and output_path.exists():
        raise partial.ReceiptRecoveryError('EXPORT_LEGACY_UNBOUND_ORPHAN','Existing workbook has no trusted pre-export intent; preserve it for explicit inspection')
    if intent and output_path.is_file():partial.validate_publication(root,intent_path,output_path)
    verdict=partial.human_state(directory)
    if not verdict['completed'] or not verdict['approved'] or batch_path!=directory/'approved_generation_batch_v1.json':
        raise partial.ReceiptRecoveryError('HUMAN_BUNDLE_INVALID','Export requires the complete native Human Approval bundle')
    approved=partial.read(batch_path);approval_path=directory/'generation_batch_approval_receipt.json'
    approval=partial.read(approval_path);source=Path(approval['source_batch']['path']).resolve()
    if recovery.ref(root,source)['sha256']!=approval['source_batch']['sha256']:
        raise partial.ReceiptRecoveryError('DELIVERY_SOURCE_CHANGED','Human Approval source changed')
    recovery.validate_approved_source(root,request_id,approved,approval)
    sources=partial.source_snapshot(root,source)
    request=partial.read(recovery.resolve_ref(root,sources['request_ref']))
    value={'schema_version':partial.EXPORT,'request_id':request_id,'business_id':request['persona_id'],
        'operation_id':intent['operation_id'] if intent else uuid.uuid4().hex,'actor':authorization_context['actor'],
        'authorization':intent['authorization'] if intent else proof,'sources':sources,
        'approved_batch_ref':recovery.ref(root,batch_path),'approval_receipt_ref':recovery.ref(root,approval_path),
        'template_ref':recovery.ref(root,template_path),'output_path':output_path.relative_to(root).as_posix(),
        'exporter_version':EXPORTER_VERSION,'mode':'mix'}
    partial.publish(intent_path,value,partial.EXPORT)
    return intent_path,value


def validate_workbook(path,template_path,contents):
    from openpyxl import load_workbook
    with zipfile.ZipFile(path) as archive:
        if sum(x.file_size for x in archive.infolist())>64*1024*1024 or archive.testzip() is not None:
            raise partial.ReceiptRecoveryError('EXPORT_PUBLICATION_INVALID','Workbook ZIP/CRC validation failed')
    book=load_workbook(path,read_only=True,data_only=False)
    template=load_workbook(template_path,read_only=True,data_only=False)
    try:
        if book.sheetnames!=template.sheetnames or book._external_links:
            raise ValueError('Unexpected workbook sheets or external links')
        sheet=book.active;source=template.active
        if (sheet.cell(4,1).value,sheet.cell(4,2).value)!=EXPECTED_HEADERS:raise ValueError('Frozen headers differ')
        for row in range(1,max(sheet.max_row,source.max_row,len(contents)+4)+1):
            for column in range(1,max(sheet.max_column,source.max_column,2)+1):
                expected=source.cell(row,column).value if row<5 else None
                if 5<=row<5+len(contents) and column in {1,2}:
                    expected=contents[row-5]['title' if column==1 else 'narration']
                if sheet.cell(row,column).value!=expected:raise ValueError('Workbook content, count or metadata differs')
        for name in book.sheetnames:
            if name!=sheet.title and list(book[name].values)!=list(template[name].values):raise ValueError('Unexpected extra sheet content')
    except (ValueError,KeyError,TypeError) as exc:
        raise partial.ReceiptRecoveryError('EXPORT_PUBLICATION_INVALID',str(exc)) from exc
    finally:book.close();template.close()


@compatibility_guarded
def export_mix_excel(
    batch_path: Path,
    template_path: Path,
    output_path: Path,
    fixture_only: bool = False,
    authorization_context: dict[str, Any] | None = None,
) -> dict[str, Any]:
    from openpyxl import load_workbook
    batch_path = batch_path.expanduser().resolve()
    template_path = template_path.expanduser().resolve()
    output_path = output_path.expanduser().resolve()
    for path in (batch_path, template_path):
        if not path.is_file():
            raise FileNotFoundError(path)
    if output_path == template_path:
        raise RuntimeError("Excel output must not overwrite the source Template.")
    batch = json.loads(batch_path.read_text(encoding="utf-8"))
    from generation_recovery_v1 import validate_approved_source, runtime_for_generation, request_root, read, RecoveryError
    root = runtime_for_generation(batch_path, batch)
    if root and (request_root(root, str(batch.get('request_id') or '')) / 'recovery').exists():
        expected = request_root(root, batch['request_id']) / 'approved_generation_batch_v1.json'
        if batch_path != expected:
            raise RecoveryError('GENERATION_RECOVERY_SOURCE_CHANGED', 'Recovery exports require the canonical Request-root approved bundle')
        validate_approved_source(root, batch['request_id'], batch,
            read(expected.parent / 'generation_batch_approval_receipt.json'))
    if batch.get("status") != "approved":
        raise RuntimeError("Excel Export requires an Approved Generation Batch.")
    if batch.get("profile") != "mix":
        raise RuntimeError("Mix Excel Export rejects non-mix batches.")
    if batch.get("human_review", {}).get("all_export_items_human_approved") is not True:
        raise RuntimeError("Every exported item must be Human Approved.")
    contents = batch.get("contents") or []
    if not contents or any(item.get("status") != "human_approved" for item in contents):
        raise RuntimeError("Approved Batch contains a non-approved export item.")
    persona_ref = (batch.get("lineage") or {}).get("persona") or {}
    if persona_ref.get("path"):
        from customer_constraints_v1 import assert_publication_allowed, publication_text
        persona_path = Path(persona_ref["path"])
        if not persona_path.is_file() or sha256_file(persona_path) != persona_ref.get("sha256"):
            raise RuntimeError("Export Business Persona lineage does not validate.")
        assert_publication_allowed(json.loads(persona_path.read_text(encoding="utf-8")), persona_path,
            [publication_text(item) for item in contents])
    intent_path=None
    if root and authorization_context is not None:
        intent_path,intent=prepare_mix_export(pipeline_root=root,request_id=batch['request_id'],batch_path=batch_path,
            template_path=template_path,output_path=output_path,authorization_context=authorization_context)
    elif root and (partial.export_paths(root,batch['request_id'])[0].exists() or
                   (root / __import__('profile_compatibility_release_v1_1').CONFIG).exists()):
        raise partial.ReceiptRecoveryError('DELIVERY_AUTHORIZATION_REQUIRED','Current trusted export authorization is required')
    if output_path.exists() and intent_path is None:
        raise RuntimeError("Excel output already exists and cannot be overwritten.")

    template_sha_before = sha256_file(template_path)
    workbook = load_workbook(template_path)
    sheet = workbook.active
    headers = (sheet.cell(4, 1).value, sheet.cell(4, 2).value)
    if headers != EXPECTED_HEADERS:
        workbook.close()
        raise RuntimeError(
            f"Mix Excel Row 4 headers mismatch: expected {EXPECTED_HEADERS!r}, got {headers!r}."
        )
    workbook.close()

    output_path.parent.mkdir(parents=True,exist_ok=True)
    staged=intent_path.parent/'mix_export_workbook_v1.xlsx' if intent_path else output_path
    if not output_path.exists():
        if not staged.exists():
            # Stage is retained as an immutable recovery artifact. A half file
            # cannot be overwritten to hide a failed persistence step.
            temporary=staged.parent/('.'+uuid.uuid4().hex+'.tmp.xlsx')
            workbook=load_workbook(template_path);sheet=workbook.active
            style_title=copy(sheet.cell(5,1)._style);style_narration=copy(sheet.cell(5,2)._style)
            for row in range(5,max(5,sheet.max_row)+1):
                for column in range(1,max(2,sheet.max_column)+1):sheet.cell(row,column).value=None
            for index,item in enumerate(contents,start=5):
                if not item.get('title') or not item.get('narration'):raise RuntimeError('Excel title and narration are required')
                sheet.cell(index,1,item['title'])._style=copy(style_title)
                sheet.cell(index,2,item['narration'])._style=copy(style_narration)
            workbook.save(temporary);workbook.close()
            with temporary.open('r+b') as handle:handle.flush();os.fsync(handle.fileno())
            validate_workbook(temporary,template_path,contents)
            os.link(temporary,staged)
            temporary.unlink()
        validate_workbook(staged,template_path,contents)
        if intent_path:
            partial.identity(authorization_context,intent)
            prepare_mix_export(pipeline_root=root,request_id=batch['request_id'],batch_path=batch_path,
                template_path=template_path,output_path=output_path,authorization_context=authorization_context)
            _,publication_path=partial.export_paths(root,batch['request_id'])
            partial.publish(publication_path,{'schema_version':partial.PUBLICATION,'intent_ref':recovery.ref(root,intent_path),
                'output_path':intent['output_path'],'output_sha256':sha256_file(staged)},partial.PUBLICATION)
            os.link(staged,output_path)  # Exclusive publication, never replace.
    if intent_path:partial.validate_publication(root,intent_path,output_path)
    validate_workbook(output_path,template_path,contents)
    template_sha_after = sha256_file(template_path)
    if template_sha_after != template_sha_before:
        raise RuntimeError("Source Excel Template was modified.")
    receipt = {
        "schema_version": partial.EXPORT_RECEIPT if intent_path else "mix-excel-export-receipt-v1.0",
        "exporter_version": EXPORTER_VERSION,
        "fixture_only": fixture_only,
        "batch_path": str(batch_path),
        "batch_sha256": sha256_file(batch_path),
        "template_path": str(template_path),
        "template_sha256": template_sha_before,
        "output_path": str(output_path),
        "output_sha256": sha256_file(output_path),
        "exported_row_count": len(contents),
        "row_4_headers": list(EXPECTED_HEADERS),
        "first_data_row": 5,
        "internal_metadata_columns_added": False,
        "source_template_modified": False,
        "validation_passed": True,
    }
    if intent_path:
        partial.identity(authorization_context,intent)
        prepare_mix_export(pipeline_root=root,request_id=batch['request_id'],batch_path=batch_path,
            template_path=template_path,output_path=output_path,authorization_context=authorization_context)
        receipt.update(export_intent_ref=recovery.ref(root,intent_path),publication_ref=recovery.ref(root,publication_path if 'publication_path' in locals() else partial.export_paths(root,batch['request_id'])[1]))
        partial.check(receipt,partial.EXPORT_RECEIPT)
    return receipt


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Export a Human Approved Mix Script Batch into the frozen Excel Contract."
    )
    parser.add_argument("--batch", required=True)
    parser.add_argument("--template", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--fixture-only", action="store_true")
    parser.add_argument('--authorization-file')
    args = parser.parse_args()
    receipt = export_mix_excel(
        Path(args.batch),
        Path(args.template),
        Path(args.output),
        args.fixture_only,
        authorization_context=partial.read(Path(args.authorization_file)) if args.authorization_file else None,
    )
    receipt_path = Path(args.output).expanduser().resolve().with_suffix(
        ".export_receipt.json"
    )
    partial.write_once(receipt_path,receipt)
    print("MIX EXCEL EXPORT V1 PASS")
    print(f"Rows: {receipt['exported_row_count']}")
    print(f"Fixture only: {receipt['fixture_only']}")
    print(f"Output: {receipt['output_path']}")
    print(f"Receipt: {receipt_path}")


if __name__ == "__main__":
    main()
