from __future__ import annotations

import hashlib
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from approve_generation_batch_v1 import approve_generation_batch  # noqa: E402
from export_mix_excel_v1 import EXPECTED_HEADERS, export_mix_excel  # noqa: E402
from generate_mix_scripts_v1 import (  # noqa: E402
    apply_metadata_fact_lineage_fix,
    assess_claim_support,
    bounded_fact_violations,
    build_content_angle_plan,
    detect_grounding_risks,
    generate_batch,
    generate_selective_revision_batch,
    get_client,
    invoke_generation_transport,
    prepare_generation_egress,
    review_flag_summary,
    render_revision_prompt,
    render_review_pack,
    safe_speaker_projection,
    semantic_diversity_summary,
    structural_fingerprint_projection,
    validate_generated_item,
    validate_generation_inputs,
    write_batch,
    write_review_pack_for_existing_batch,
)


def write_json(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class ProductionMvpPhase1Tests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.persona_path = self.root / "persona.json"
        self.request_path = self.root / "request.json"
        self.pattern_path = self.root / "pattern.json"
        self.plan_path = self.root / "plan.json"
        self.fingerprint_paths: list[Path] = []
        persona = self.persona()
        write_json(self.persona_path, persona)
        request = self.request()
        write_json(self.request_path, request)
        pattern = self.pattern()
        write_json(self.pattern_path, pattern)
        pool = []
        for index, case_id in enumerate(("100", "200", "300"), start=1):
            path = self.root / f"fingerprint_{case_id}.json"
            write_json(path, self.fingerprint(case_id, index))
            self.fingerprint_paths.append(path)
            pool.append(
                {
                    "case_id": case_id,
                    "approved_case_sha": hashlib.sha256(case_id.encode()).hexdigest(),
                    "fingerprint_sha": sha256(path),
                    "ranking_score": 1.0 - index / 10,
                }
            )
        write_json(
            self.plan_path,
            {
                "request_id": request["request_id"],
                "persona": {"persona_sha": sha256(self.persona_path)},
                "request": {"request_sha": sha256(self.request_path)},
                "coverage": {"status": "supported"},
                "selected_patterns": [
                    {
                        "pattern_id": pattern["pattern_id"],
                        "approved_pattern_sha": sha256(self.pattern_path),
                    }
                ],
                "eligible_case_pool": pool,
                "rotation": {
                    "case_rotation_order": ["100", "200", "300"]
                },
                "excluded_sources": [
                    {
                        "source_id": "rejected_pattern",
                        "source_type": "pattern_candidate",
                    }
                ],
            },
        )

    def tearDown(self) -> None:
        self.temp.cleanup()

    @staticmethod
    def fact(value, state="known") -> dict:
        return {"state": state, "value": value}

    def persona(self) -> dict:
        return {
            "persona_id": "fixture_persona",
            "revision": 1,
            "fixture_only": True,
            "lifecycle": {"status": "approved", "approved": True},
            "facts": {
                "company_short_name": self.fact("Fixture Green Service"),
                "industry": self.fact("绿植租赁"),
                "years_in_business": self.fact(None, "unknown"),
                "service_area": self.fact(None, "unknown"),
                "business_address": self.fact(None, "unknown"),
                "primary_products_or_services": self.fact(["绿植租赁", "上门养护"]),
                "core_audience": self.fact(["企业客户"]),
                "customer_pains": self.fact(["缺少养护能力"]),
                "differentiators": self.fact(["现场勘测"]),
                "brand_story": self.fact("认真做好空间绿植服务"),
                "process_facts": self.fact(["勘测", "布置", "养护"]),
                "service_process": self.fact(["了解需求", "现场勘测", "持续养护"]),
                "values": self.fact(["清晰沟通"]),
                "unknown_facts": self.fact(["经营年限", "复购率"]),
                "forbidden_claims": self.fact(["效果保证"]),
                "prohibited_topics": self.fact(["虚构客户反馈"]),
            },
            "provenance": {"content_sha256": "fixture-content-sha"},
        }

    @staticmethod
    def request() -> dict:
        return {
            "request_id": "fixture_request",
            "profile": "mix",
            "quantity": 10,
            "platform": "douyin",
            "content_intent": "mixed",
            "cta_intent": "optional",
            "constraints": {"generation_must_not_start": False},
        }

    @staticmethod
    def pattern() -> dict:
        return {
            "pattern_id": "approved_pattern",
            "status": "approved",
            "working_name": "Narration-led Process Projection",
            "definition": {
                "invariants": [
                    "Narration is the semantic spine.",
                    "Real process scenes provide concrete visual projection.",
                ]
            },
            "scope": {"current_generalization_limit": "fixture"},
            "effectiveness": {"status": "unvalidated"},
        }

    @staticmethod
    def fingerprint(case_id: str, index: int) -> dict:
        return {
            "case_id": case_id,
            "narration_features": {
                "segment_count": index + 2,
                "speech_to_video_ratio": 0.8,
                "source_narration": "经营二十一年",
            },
            "visual_shot_features": {
                "shot_count": 10 + index,
                "shot_duration_stats": {"mean": 1.5},
                "primary_role_sequence": ["hook", "action", "product"],
                "primary_role_counts": {"hook": 1, "action": 6, "product": 3},
                "ocr_raw": "某客户原始字幕",
            },
            "structure_features": {
                "structure_signature": "hook>development>payoff",
                "structure_sequence": [
                    {"stage": "hook", "description": "某店客户很多"},
                    {"stage": "payoff", "description": "某城市最好"},
                ],
                "hook_modalities": ["audio", "visual_scene"],
                "hook_candidate": {"audio": "客户都说好"},
            },
            "audio_visual_features": {
                "scene_relation_counts": {"supplement": 10},
                "shot_to_narration_cardinality": "many_to_many",
            },
            "content_features": {
                "claims": ["经营二十一年", "客户都认可"],
                "content_goal_candidate": "某城市商家主张",
            },
        }

    def context(self) -> dict:
        return validate_generation_inputs(
            self.persona_path,
            self.request_path,
            self.plan_path,
            self.pattern_path,
            self.fingerprint_paths,
        )

    @staticmethod
    def model_items(count: int = 10) -> list[dict]:
        angles = [
            "空间需求先了解",
            "现场勘测看细节",
            "方案确认再布置",
            "绿植选择看场景",
            "日常养护有过程",
            "办公空间慢慢调整",
            "服务从沟通开始",
            "布置之后持续照料",
            "不同空间不同考虑",
            "把服务过程讲清楚",
        ]
        narrations = [
            "做办公空间绿植服务，先把使用需求听清楚，再去现场看光线和空间，让后面的布置有具体依据。",
            "绿植租赁不只是把植物搬进办公室，现场勘测能帮助我们看清空间条件，再讨论合适的布置方向。",
            "从了解需求到确认方案，每一步都需要沟通。画面可以记录勘测和讨论，让服务过程更具体。",
            "同样是办公空间，使用方式可能不同。我们会先看现场，再结合需求考虑绿植布置和后续养护。",
            "上门养护是持续服务的一部分。检查植物状态、整理叶片和调整摆放，都是可以被记录的日常过程。",
            "空间里的绿植不是一次布置就结束，后续还需要结合实际状态持续照料，让服务过程保持清晰。",
            "我们更愿意先把需求问明白，再进入现场勘测和方案沟通。真实过程比空泛表达更容易理解。",
            "布置完成之后，定期养护会继续发生。镜头记录工作人员的实际动作，也补充了旁白之外的细节。",
            "面对不同办公空间，我们从现场条件出发，讨论布置与养护安排，而不是先给出没有依据的结论。",
            "把了解需求、现场勘测、实施布置和后续养护依次讲清楚，观众才能看到服务具体发生在哪里。",
        ]
        return [
            {
                "slot_id": f"SLOT_{index:03d}",
                "title": angles[index - 1],
                "narration": narrations[index - 1],
                "persona_fact_refs_used": [
                    "primary_products_or_services",
                    "process_facts",
                    "service_process",
                ],
                "claim_candidates": [],
                "cta": {"present": False, "text": None},
            }
            for index in range(1, count + 1)
        ]

    def transport(self):
        calls = {"count": 0}

        def call(_prompt: str) -> dict:
            calls["count"] += 1
            payload = {"items": self.model_items()}
            return {
                "payload": payload,
                "response_sha256": hashlib.sha256(json.dumps(payload).encode()).hexdigest(),
                "usage": {"prompt_tokens": 100, "completion_tokens": 200, "total_tokens": 300},
                "elapsed_seconds": 0.1,
            }

        return calls, call

    def generated_batch(self) -> dict:
        _calls, transport = self.transport()
        return generate_batch(
            self.context(),
            transport=transport,
            created_at="2026-09-11T00:00:00+00:00",
        )

    def test_unapproved_persona_is_rejected(self) -> None:
        persona = json.loads(self.persona_path.read_text(encoding="utf-8"))
        persona["lifecycle"] = {"status": "review_required", "approved": False}
        write_json(self.persona_path, persona)
        with self.assertRaisesRegex(RuntimeError, "Approved Persona"):
            self.context()

    def test_unsupported_plan_and_news_request_are_rejected(self) -> None:
        plan = json.loads(self.plan_path.read_text(encoding="utf-8"))
        plan["coverage"] = {"status": "insufficient"}
        write_json(self.plan_path, plan)
        with self.assertRaisesRegex(RuntimeError, "coverage"):
            self.context()
        plan["coverage"] = {"status": "supported"}
        write_json(self.plan_path, plan)
        request = json.loads(self.request_path.read_text(encoding="utf-8"))
        request["profile"] = "news"
        write_json(self.request_path, request)
        plan["request"]["request_sha"] = sha256(self.request_path)
        write_json(self.plan_path, plan)
        with self.assertRaisesRegex(RuntimeError, "non-mix"):
            self.context()

    def test_rejected_pattern_is_rejected(self) -> None:
        pattern = json.loads(self.pattern_path.read_text(encoding="utf-8"))
        pattern["status"] = "review_rejected"
        write_json(self.pattern_path, pattern)
        plan = json.loads(self.plan_path.read_text(encoding="utf-8"))
        plan["selected_patterns"][0]["approved_pattern_sha"] = sha256(self.pattern_path)
        write_json(self.plan_path, plan)
        with self.assertRaisesRegex(RuntimeError, "Approved Pattern"):
            self.context()

    def test_case_raw_text_and_facts_are_not_in_structural_projection_or_prompt(self) -> None:
        fingerprint = self.fingerprint("100", 1)
        projection = structural_fingerprint_projection(fingerprint, "CASE_REF_001")
        serialized = json.dumps(projection, ensure_ascii=False)
        for forbidden in ("经营二十一年", "某客户原始字幕", "客户都认可", "某城市最好"):
            self.assertNotIn(forbidden, serialized)
        context = self.context()
        _safe, prompt, _audit = prepare_generation_egress(
            context,
            [{"slot_id": "SLOT_001", "case_id": "100"}],
            [],
        )
        self.assertNotIn("经营二十一年", prompt)
        self.assertNotIn("某客户原始字幕", prompt)
        self.assertNotIn("100", prompt)

    def test_unknown_fact_cannot_be_sent_or_referenced(self) -> None:
        context = self.context()
        safe, _prompt, _audit = prepare_generation_egress(
            context,
            [{"slot_id": "SLOT_001", "case_id": "100"}],
            [],
        )
        self.assertNotIn("years_in_business", safe["persona"]["facts"])
        item = self.model_items(1)[0]
        item["persona_fact_refs_used"] = ["years_in_business"]
        valid, errors, _flags = validate_generated_item(
            item,
            set(safe["persona"]["facts"]),
            safe["persona"],
            [],
            "optional",
        )
        self.assertIsNone(valid)
        self.assertTrue(any("invalid_or_non_known" in error for error in errors))

    def test_unsupported_numeric_fact_and_process_proof_are_rejected(self) -> None:
        context = self.context()
        safe, _prompt, _audit = prepare_generation_egress(
            context,
            [{"slot_id": "SLOT_001", "case_id": "100"}],
            [],
        )
        item = self.model_items(1)[0]
        item["narration"] = "我们经营二十一年，这证明客户都认可。"
        valid, errors, _flags = validate_generated_item(
            item,
            set(safe["persona"]["facts"]),
            safe["persona"],
            [],
            "optional",
        )
        self.assertIsNone(valid)
        self.assertTrue(any("unsupported_numeric_fact" in error for error in errors))
        self.assertTrue(any("upgraded_to_proof" in error for error in errors))

    def test_batch_duplicate_is_detected(self) -> None:
        context = self.context()
        safe, _prompt, _audit = prepare_generation_egress(
            context,
            [{"slot_id": "SLOT_001", "case_id": "100"}],
            [],
        )
        item = self.model_items(1)[0]
        previous = {
            "title": item["title"],
            "narration": item["narration"],
        }
        valid, errors, _flags = validate_generated_item(
            item,
            set(safe["persona"]["facts"]),
            safe["persona"],
            [previous],
            "optional",
        )
        self.assertIsNone(valid)
        self.assertTrue(any("duplicate" in error for error in errors))

    def test_quantity_ten_batch_is_review_required_with_lineage(self) -> None:
        batch = self.generated_batch()
        self.assertEqual(batch["machine_pass_count"], 10)
        self.assertEqual(batch["status"], "review_required")
        self.assertEqual(len(batch["contents"]), 10)
        self.assertTrue(batch["validation"]["sha_lineage_complete"])
        self.assertEqual(batch["lineage"]["persona"]["sha256"], sha256(self.persona_path))
        self.assertEqual(batch["lineage"]["request"]["sha256"], sha256(self.request_path))
        self.assertEqual(batch["lineage"]["source_plan"]["sha256"], sha256(self.plan_path))
        self.assertEqual(batch["lineage"]["approved_pattern"]["sha256"], sha256(self.pattern_path))

    def test_content_angle_plan_has_controlled_unique_slots_and_capacity(self) -> None:
        plan = build_content_angle_plan(self.context())
        self.assertEqual(plan["planned_quantity"], 10)
        self.assertGreaterEqual(plan["high_confidence_distinct_capacity"], 10)
        self.assertEqual(len(plan["slots"]), 10)
        self.assertEqual(
            len({slot["primary_angle"] for slot in plan["slots"]}), 10
        )
        for slot in plan["slots"]:
            self.assertTrue(slot["primary_persona_fact_refs"])
            self.assertTrue(slot["opening_strategy"])
            self.assertTrue(slot["narrative_mode"])
            self.assertTrue(slot["selected_case_structural_reference"])

    def test_review_flags_aggregate_from_all_items(self) -> None:
        contents = [
            {
                "content_id": f"C{index:03d}",
                "potential_review_flags": [
                    "claim_candidates_require_human_review"
                ],
            }
            for index in range(1, 11)
        ]
        summary = review_flag_summary(contents)
        self.assertEqual(summary["items_with_review_flags"], 10)
        self.assertEqual(summary["total_review_flag_count"], 10)
        self.assertEqual(summary["claim_review_flag_count"], 10)

    def test_claim_language_and_semantic_diversity_are_audited(self) -> None:
        context = self.context()
        slot = build_content_angle_plan(context)["slots"][0]
        safe, _prompt, _audit = prepare_generation_egress(context, [slot], [])
        item = self.model_items(1)[0]
        item["central_claim"] = "现场勘测才是关键"
        item["narration"] = "办公绿植服务里，现场勘测才是关键。"
        valid, _errors, flags = validate_generated_item(
            item,
            set(safe["persona"]["facts"]),
            safe["persona"],
            [],
            "optional",
            slot,
        )
        self.assertIsNotNone(valid)
        self.assertIn(
            "claim_language_requires_human_review:才是关键", flags
        )
        contents = [
            {
                "content_id": "C001",
                "primary_angle": "pain",
                "central_claim": "现场勘测先看空间需求",
                "narration": "了解需求，现场勘测，再确认方案。",
                "persona_fact_refs_used": ["customer_pains"],
            },
            {
                "content_id": "C002",
                "primary_angle": "pain",
                "central_claim": "现场勘测要先看空间需求",
                "narration": "了解需求，现场勘测，再确认方案。",
                "persona_fact_refs_used": ["customer_pains"],
            },
        ]
        diversity = semantic_diversity_summary(contents)
        self.assertIn("primary_angle_reuse", diversity["semantic_diversity_flags"])
        self.assertIn("central_claim_overlap", diversity["semantic_diversity_flags"])
        self.assertIn(
            "service_process_sequence_repetition",
            diversity["semantic_diversity_flags"],
        )

    def test_privacy_gate_runs_before_transport_and_pii_failure_calls_zero(self) -> None:
        calls = {"count": 0}

        def transport(_prompt: str) -> dict:
            calls["count"] += 1
            return {}

        unsafe = "收件人：张三，身份证：11010519491231002X"
        with self.assertRaises(RuntimeError):
            invoke_generation_transport(
                safe_input={},
                rendered_prompt=unsafe,
                pre_call_audit={
                    "runtime_persona_projection_sha256": "x",
                },
                transport=transport,
            )
        self.assertEqual(calls["count"], 0)

    def test_frontline_speaker_authority_and_real_customer_qualifiers(self) -> None:
        speaker = {
            "persona_id": "frontline",
            "persona_scope": "speaker",
            "speaker_type": "frontline_expert",
            "facts": {
                "public_display_name": self.fact("林东方"),
                "public_role": self.fact("主厨"),
                "speaker_role_facts": self.fact(["负责代炒菜窗口的烹饪工作"]),
                "first_person_allowed_topics": self.fact(["现场烹饪流程"]),
                "first_person_forbidden_claims": self.fact(["我是老板", "我决定免租"]),
                "unknown_facts": self.fact(["从业年限", "是否创始人"]),
            },
        }
        safe_speaker = safe_speaker_projection(speaker)
        item = self.model_items(1)[0]
        item["narration"] = "我是老板，我创办了这家食堂。"
        item["speaker_fact_refs_used"] = ["speaker_role_facts"]
        valid, errors, _flags = validate_generated_item(
            item,
            {"customer_pains"},
            {"facts": {"customer_pains": ["不太会做饭"]}, "guardrails": {}},
            [],
            "none",
            allowed_speaker_fact_refs={"speaker_role_facts"},
            safe_speaker=safe_speaker,
        )
        self.assertIsNone(valid)
        self.assertTrue(any("frontline_speaker_authority_exceeded" in error for error in errors))

        safe_business = {
            "facts": {
                "pricing_facts": ["素菜加工费通常 8-10 元", "清蒸约 15 元", "红烧约 18 元左右"],
                "service_time_facts": ["部分快速菜品约 5 分钟", "鱼类等较复杂菜品约一刻钟", "高峰期可能排队"],
                "time_efficiency_fact": ["从购买食材到菜品上桌，全程约 30 分钟"],
            }
        }
        violations = bounded_fact_violations(
            "素菜加工费固定8元，任何菜30分钟搞定，而且不用等。",
            safe_business,
        )
        self.assertIn("vegetable_price_range_or_qualifier_lost", violations)
        self.assertIn("thirty_minute_scope_or_qualifier_lost", violations)
        self.assertIn("waiting_limit_removed", violations)

    def test_public_location_guard_does_not_treat_market_terms_as_locations(self) -> None:
        safe = {
            "facts": {
                "customer_use_cases": ["顾客在菜市场购买食材后送到窗口加工"],
                "location_public_area": "上海浦东新区三林镇",
            },
            "guardrails": {},
        }
        base = self.model_items(1)[0]
        base["persona_fact_refs_used"] = ["customer_use_cases"]
        base["narration"] = "顾客在菜市场买完食材，可以送到社区食堂窗口现场加工。"
        valid, errors, _flags = validate_generated_item(
            base,
            {"customer_use_cases", "location_public_area"},
            safe,
            [],
            "optional",
        )
        self.assertIsNotNone(valid, errors)

        invented = dict(base)
        invented["narration"] = "这项服务位于北京市朝阳区。"
        valid, errors, _flags = validate_generated_item(
            invented,
            {"customer_use_cases", "location_public_area"},
            safe,
            [],
            "optional",
        )
        self.assertIsNone(valid)
        self.assertIn("unsupported_location_fact", errors)

    def test_known_price_numbers_allow_spacing_but_keep_qualifier_guards(self) -> None:
        safe = {
            "facts": {
                "pricing_facts": [
                    "素菜加工费通常 8-10 元",
                    "清蒸约 15 元",
                    "红烧约 18 元左右",
                ]
            },
            "guardrails": {},
        }
        item = self.model_items(1)[0]
        item["persona_fact_refs_used"] = ["pricing_facts"]
        item["narration"] = "素菜加工费通常8到10元，清蒸约15元，红烧约18元左右。"
        valid, errors, _flags = validate_generated_item(
            item,
            {"pricing_facts"},
            safe,
            [],
            "optional",
        )
        self.assertIsNotNone(valid, errors)

    def test_metadata_fact_lineage_fix_preserves_content_text(self) -> None:
        source = self.model_items(1)[0]
        source["central_claim"] = "价格公开清晰"
        source["narration"] += "价格公开清晰。"
        source["persona_fact_refs_used"] = ["pricing_facts"]
        before = {
            field: source.get(field)
            for field in ("title", "narration", "central_claim", "cta")
        }
        revised = apply_metadata_fact_lineage_fix(
            source,
            {
                "add_persona_fact_refs_used": ["differentiators"],
                "append_claim_candidates": [
                    {
                        "text": "价格公开清晰",
                        "persona_fact_refs": ["differentiators"],
                        "speaker_fact_refs": [],
                    }
                ],
            },
            {"pricing_facts", "differentiators"},
            "a" * 64,
            "b" * 64,
            2,
        )
        after = {
            field: revised.get(field)
            for field in ("title", "narration", "central_claim", "cta")
        }
        self.assertEqual(before, after)
        self.assertIn("differentiators", revised["persona_fact_refs_used"])
        self.assertEqual(
            revised["revision_lineage"]["source_text_sha256"],
            revised["revision_lineage"]["revised_text_sha256"],
        )
        with self.assertRaisesRegex(RuntimeError, "already exist"):
            apply_metadata_fact_lineage_fix(
                source,
                {
                    "append_claim_candidates": [
                        {
                            "text": "不存在的新主张",
                            "persona_fact_refs": ["differentiators"],
                            "speaker_fact_refs": [],
                        }
                    ]
                },
                {"differentiators"},
                "a" * 64,
                "b" * 64,
                2,
            )

    def test_batch_write_creates_human_review_pack(self) -> None:
        batch = self.generated_batch()
        batch_path, review_path, _digest = write_batch(batch, self.root / "batches")
        self.assertTrue(batch_path.is_file())
        review_text = review_path.read_text(encoding="utf-8")
        self.assertIn(batch["contents"][0]["title"], review_text)
        self.assertIn("Human review is required", review_text)
        self.assertIn("Source Batch SHA-256", review_text)
        self.assertIn("### Supporting Persona Facts", review_text)

    def test_review_pack_shows_verbatim_safe_known_facts_and_claim_mapping(self) -> None:
        batch = self.generated_batch()
        item = batch["contents"][0]
        item["claim_candidates"] = [
            {
                "claim": "缺少养护能力",
                "status": "candidate_unverified",
                "persona_fact_refs": ["customer_pains"],
            }
        ]
        text = render_review_pack(batch, "a" * 64)
        self.assertIn("缺少养护能力", text)
        self.assertIn("Supporting Fact Refs: `customer_pains`", text)
        self.assertIn("Support Status: `direct`", text)
        self.assertNotIn("经营年限", text)

    def test_review_pack_marks_non_known_and_redacts_persona_pii(self) -> None:
        persona = json.loads(self.persona_path.read_text(encoding="utf-8"))
        persona["facts"]["customer_pains"] = self.fact(
            ["联系人：张三，手机号：13812345678"]
        )
        persona["facts"]["review_fact"] = self.fact(
            "待人工确认的私人资料", "requires_review"
        )
        persona["facts"]["unknown_fact"] = self.fact(None, "unknown")
        write_json(self.persona_path, persona)
        plan = json.loads(self.plan_path.read_text(encoding="utf-8"))
        plan["persona"]["persona_sha"] = sha256(self.persona_path)
        write_json(self.plan_path, plan)
        batch = self.generated_batch()
        batch["lineage"]["persona"]["sha256"] = sha256(self.persona_path)
        batch["contents"][0]["persona_fact_refs_used"].extend(
            ["customer_pains", "review_fact", "unknown_fact"]
        )
        text = render_review_pack(batch, "b" * 64)
        self.assertNotIn("张三", text)
        self.assertNotIn("13812345678", text)
        self.assertIn("[姓名]", text)
        self.assertIn("[REQUIRES_REVIEW: not eligible as supporting fact]", text)
        self.assertIn("[UNKNOWN: not eligible as supporting fact]", text)
        self.assertNotIn("待人工确认的私人资料", text)

    def test_review_pack_only_keeps_batch_sha_and_calls_no_remote_model(self) -> None:
        batch = self.generated_batch()
        batch_path, review_path, batch_sha = write_batch(batch, self.root / "review")
        with patch("generate_mix_scripts_v1.get_client") as client:
            output, after_sha, projection = write_review_pack_for_existing_batch(
                batch_path, review_path
            )
        self.assertEqual(output, review_path)
        self.assertEqual(after_sha, batch_sha)
        self.assertEqual(sha256(batch_path), batch_sha)
        self.assertFalse(projection["remote_model_used"])
        client.assert_not_called()

    def test_first_person_and_inference_expansion_are_not_direct(self) -> None:
        known = {"values": ["清晰沟通"], "customer_pains": ["缺少养护能力"]}
        first_person = assess_claim_support(
            "作为经营者，我经常看到绿植需要持续跟进",
            ["values"],
            known,
        )
        inference = assess_claim_support(
            "很多人往往都是因为缺少养护能力",
            ["customer_pains"],
            known,
        )
        self.assertNotEqual(first_person["status"], "direct")
        self.assertNotEqual(inference["status"], "direct")

    def test_grounding_rules_distinguish_questions_and_unsupported_claims(self) -> None:
        safe = {
            "facts": {
                "customer_pains": ["空间绿植状态不稳定"],
                "primary_products_or_services": ["绿植租赁", "定期养护"],
                "founder_or_operator_story": ["经营者参与现场勘测和日常服务"],
            }
        }
        question = detect_grounding_risks(
            "办公室绿植状态不稳定，可以先看什么？",
            ["customer_pains"],
            safe,
        )
        unsupported = detect_grounding_risks(
            "不少企业会遇到这个问题，有人会以为需要反复调整布置。有人以为布置就是终点，客户会提到养护问题。很多人往往认为状态常常和养护有关，让状态更稳定，也比单次布置更省心。",
            ["customer_pains"],
            safe,
        )
        bundle = detect_grounding_risks(
            "绿植租赁包含定期养护。",
            ["primary_products_or_services"],
            safe,
        )
        supported_operator = detect_grounding_risks(
            "我会参与现场勘测和日常服务。",
            ["founder_or_operator_story"],
            safe,
        )
        softened_causality = detect_grounding_risks(
            "空间绿植状态不稳定，可能与缺少持续养护能力有关。",
            ["customer_pains"],
            safe,
        )
        self.assertEqual(question, [])
        self.assertTrue(any("empirical_frequency_claim" in flag for flag in unsupported))
        self.assertTrue(any("causal_effect_claim" in flag for flag in unsupported))
        self.assertTrue(any("assumed_customer_behavior" in flag for flag in unsupported))
        self.assertTrue(
            any("causal_effect_claim" in flag for flag in softened_causality)
        )
        self.assertTrue(any("service_bundle_expansion" in flag for flag in bundle))
        self.assertEqual(supported_operator, [])

    def test_revision_prompt_adds_pain_slot_grounding_constraints(self) -> None:
        prompt = render_revision_prompt(
            {
                "persona": {},
                "request": {},
                "approved_pattern": {},
                "case_structures": {"CASE_TEST": {}},
            },
            [
                {
                    "slot_id": "SLOT_001",
                    "case_id": "CASE_TEST",
                    "primary_angle": "pain",
                    "primary_persona_fact_refs": ["customer_pains"],
                    "optional_secondary_fact_refs": ["differentiators"],
                    "opening_strategy": "question",
                    "narrative_mode": "narration_led",
                    "selected_case_structural_reference": "CASE_TEST_STRUCTURE",
                    "original_generated_item": {},
                    "revision_instructions": "Keep the approved pain angle.",
                },
                {
                    "slot_id": "SLOT_002",
                    "case_id": "CASE_TEST",
                    "primary_angle": "misconception",
                    "primary_persona_fact_refs": ["primary_products_or_services"],
                    "optional_secondary_fact_refs": ["differentiators"],
                    "opening_strategy": "question",
                    "narrative_mode": "narration_led",
                    "selected_case_structural_reference": "CASE_TEST_STRUCTURE",
                    "original_generated_item": {},
                    "revision_instructions": "Keep the misconception angle.",
                },
            ],
            [],
        )
        self.assertIn("deterministic_slot_constraints", prompt)
        self.assertIn("State each approved pain and service fact independently", prompt)
        self.assertIn("Express the misconception as a non-empirical question", prompt)

    def test_selective_revision_preserves_approved_items_and_angles(self) -> None:
        source_batch = self.generated_batch()
        source_path, _pack, source_sha = write_batch(source_batch, self.root / "source")
        review_path = source_path.parent / "content_review.json"
        revised_ids = {
            item["content_id"]
            for item in source_batch["contents"]
            if item["content_id"].endswith(("C001", "C002", "C003", "C004", "C005", "C007", "C008", "C010"))
        }
        write_json(
            review_path,
            {
                "schema_version": "generation-batch-content-review-v1.0",
                "request_id": source_batch["request_id"],
                "source_batch_sha256": source_sha,
                "reviewer": "Fixture Reviewer",
                "decision": "partial_revision_required",
                "items": [
                    {
                        "content_id": item["content_id"],
                        "decision": "revise" if item["content_id"] in revised_ids else "approved",
                        "revision_instructions": "Use neutral KNOWN facts only."
                        if item["content_id"] in revised_ids
                        else None,
                    }
                    for item in source_batch["contents"]
                ],
            },
        )
        calls = {"count": 0}
        revise_slots = ["SLOT_001", "SLOT_002", "SLOT_003", "SLOT_004", "SLOT_005", "SLOT_007", "SLOT_008", "SLOT_010"]
        revised_payloads = [
            ("空间绿植状态怎么描述", "办公空间绿植状态不稳定。当前提供上门养护服务。", "当前提供上门养护服务。", ["customer_pains", "primary_products_or_services"]),
            ("布置之后还有哪些服务", "绿植布置不是流程最后一步。当前还提供持续上门养护。", "当前还提供持续上门养护。", ["primary_products_or_services", "service_process"]),
            ("选择服务时看哪些信息", "服务过程清晰，养护安排可追踪。", "服务过程清晰，养护安排可追踪。", ["differentiators"]),
            ("办公绿植服务流程", "服务包括了解需求、现场勘测、确认方案、实施布置和后续养护。", "当前服务有明确流程。", ["process_facts", "service_process"]),
            ("除了租赁还有上门养护", "我们提供绿植租赁，也提供持续上门养护。", "当前提供两项独立服务。", ["primary_products_or_services"]),
            ("绿植状态不稳定先看什么", "办公室绿植状态不稳定，可以先看有没有持续养护安排？", "问题聚焦持续养护安排。", ["customer_pains", "service_process"]),
            ("经营者重视的服务方式", "我们重视清晰沟通，也会持续做好服务。", "清晰沟通是当前价值表达。", ["values"]),
            ("绿植租赁与定期养护", "除了绿植租赁，我们还提供定期养护。", "当前分别提供租赁和养护服务。", ["primary_products_or_services"]),
        ]

        def transport(_prompt: str) -> dict:
            calls["count"] += 1
            items = [
                {
                    "slot_id": slot_id,
                    "title": payload[0],
                    "narration": payload[1],
                    "central_claim": payload[2],
                    "persona_fact_refs_used": payload[3],
                    "claim_candidates": [],
                    "cta": {"present": False, "text": None},
                }
                for slot_id, payload in zip(revise_slots, revised_payloads)
            ]
            payload = {"items": items}
            return {
                "payload": payload,
                "response_sha256": hashlib.sha256(json.dumps(payload).encode()).hexdigest(),
                "usage": {"prompt_tokens": 100, "completion_tokens": 200, "total_tokens": 300},
                "elapsed_seconds": 0.1,
            }

        revised = generate_selective_revision_batch(
            self.context(),
            source_path,
            review_path,
            transport=transport,
            created_at="2026-09-11T00:20:00+00:00",
        )
        original = {item["content_id"]: item for item in source_batch["contents"]}
        result = {item["content_id"]: item for item in revised["contents"]}
        self.assertEqual(calls["count"], 1)
        self.assertEqual(len(result), 10)
        self.assertEqual(result["fixture_request-C006"]["title"], original["fixture_request-C006"]["title"])
        self.assertEqual(result["fixture_request-C009"]["narration"], original["fixture_request-C009"]["narration"])
        self.assertTrue(revised["validation"]["approved_items_unchanged"])
        self.assertTrue(revised["validation"]["primary_angles_preserved"])
        self.assertEqual(revised["status"], "review_required")

        incomplete = json.loads(json.dumps(revised))
        incomplete["status"] = "generation_incomplete"
        incomplete["contents"] = [
            item
            for item in incomplete["contents"]
            if item["content_id"] != "fixture_request-C001"
        ]
        incomplete["lineage"]["source_generation_batch"] = {
            "path": str(source_path),
            "sha256": source_sha,
        }
        incomplete_path = self.root / "incomplete_revision.json"
        write_json(incomplete_path, incomplete)
        topup_calls = {"count": 0}

        def topup_transport(_prompt: str) -> dict:
            topup_calls["count"] += 1
            item = {
                "slot_id": "SLOT_001",
                "title": "空间绿植状态怎么描述",
                "narration": "办公空间绿植状态不稳定。当前提供持续上门养护。",
                "central_claim": "当前提供持续上门养护。",
                "persona_fact_refs_used": ["customer_pains", "primary_products_or_services"],
                "claim_candidates": [],
                "cta": {"present": False, "text": None},
            }
            payload = {"items": [item]}
            return {
                "payload": payload,
                "response_sha256": hashlib.sha256(json.dumps(payload).encode()).hexdigest(),
                "usage": {"prompt_tokens": 50, "completion_tokens": 50, "total_tokens": 100},
                "elapsed_seconds": 0.05,
            }

        topped_up = generate_selective_revision_batch(
            self.context(),
            source_path,
            review_path,
            batch_revision=3,
            transport=topup_transport,
            created_at="2026-09-11T00:30:00+00:00",
            carried_revision_path=incomplete_path,
        )
        self.assertEqual(topup_calls["count"], 1)
        self.assertEqual(len(topped_up["contents"]), 10)
        self.assertTrue(topped_up["authority"]["selective_revision_resume"])
        self.assertEqual(topped_up["human_content_review"]["topup_generated_count"], 1)

    def approve_batch_fixture(self) -> tuple[Path, dict]:
        batch = self.generated_batch()
        batch_path, _pack, batch_sha = write_batch(batch, self.root / "batches")
        review_path = batch_path.parent / "review.json"
        write_json(
            review_path,
            {
                "schema_version": "generation-batch-review-v1.0",
                "request_id": batch["request_id"],
                "batch_sha256": batch_sha,
                "reviewer": "Fixture Reviewer",
                "decision": "approve_items",
                "items": [
                    {"content_id": item["content_id"], "decision": "approved"}
                    for item in batch["contents"]
                ],
            },
        )
        approved_path, _receipt_path, approved, _receipt = approve_generation_batch(
            batch_path,
            review_path,
            "2026-09-11T00:10:00+00:00",
        )
        return approved_path, approved

    def test_rejected_batch_preserves_detailed_human_findings(self) -> None:
        batch = self.generated_batch()
        batch_path, _pack, batch_sha = write_batch(batch, self.root / "rejected")
        review_path = batch_path.parent / "review.json"
        findings = ["keep", "revision_required"] + ["rejected"] * 8
        write_json(
            review_path,
            {
                "schema_version": "generation-batch-review-v1.0",
                "request_id": batch["request_id"],
                "batch_sha256": batch_sha,
                "reviewer": "Fixture Reviewer",
                "decision": "reject_batch",
                "items": [
                    {"content_id": item["content_id"], "decision": finding}
                    for item, finding in zip(batch["contents"], findings)
                ],
            },
        )
        output, _receipt_path, reviewed, _receipt = approve_generation_batch(
            batch_path, review_path, "2026-09-11T00:10:00+00:00"
        )
        self.assertEqual(output.name, "reviewed_generation_batch_v1.json")
        self.assertEqual(reviewed["status"], "rejected")
        self.assertEqual(reviewed["human_review"]["finding_counts"]["keep"], 1)
        self.assertFalse(reviewed["authority"]["export_allowed"])

    def create_template(self) -> Path:
        path = self.root / "fixture_only_mix_template.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet["A1"] = "视频制作文案批量导入"
        sheet["A2"] = "填写说明"
        sheet["A3"] = "示例提示"
        sheet["A4"] = EXPECTED_HEADERS[0]
        sheet["B4"] = EXPECTED_HEADERS[1]
        sheet["A5"] = "示例标题"
        sheet["B5"] = "示例口播"
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 80
        workbook.save(path)
        workbook.close()
        return path

    def test_unapproved_batch_cannot_export(self) -> None:
        batch = self.generated_batch()
        batch_path, _pack, _digest = write_batch(batch, self.root / "batches")
        with self.assertRaisesRegex(RuntimeError, "Approved Generation Batch"):
            export_mix_excel(
                batch_path,
                self.create_template(),
                self.root / "output.xlsx",
                fixture_only=True,
            )

    def test_missing_api_key_fails_fast_without_interactive_prompt(self) -> None:
        with patch.dict(os.environ, {"DEEPSEEK_API_KEY": ""}):
            with self.assertRaisesRegex(RuntimeError, "DEEPSEEK_API_KEY is required"):
                get_client()

    def test_approved_batch_exports_frozen_excel_contract(self) -> None:
        approved_path, approved = self.approve_batch_fixture()
        template = self.create_template()
        template_sha = sha256(template)
        output = self.root / "output.xlsx"
        receipt = export_mix_excel(
            approved_path, template, output, fixture_only=True
        )
        self.assertTrue(receipt["validation_passed"])
        self.assertTrue(receipt["fixture_only"])
        self.assertEqual(receipt["exported_row_count"], 10)
        self.assertEqual(sha256(template), template_sha)
        workbook = load_workbook(output, read_only=True)
        sheet = workbook.active
        self.assertEqual((sheet["A4"].value, sheet["B4"].value), EXPECTED_HEADERS)
        self.assertEqual(sheet["A5"].value, approved["contents"][0]["title"])
        self.assertEqual(sheet["B5"].value, approved["contents"][0]["narration"])
        self.assertIsNone(sheet["C5"].value)
        serialized_cells = "\n".join(
            str(sheet.cell(row, column).value or "")
            for row in range(5, 15)
            for column in range(1, 3)
        )
        self.assertNotIn("pattern_id", serialized_cells)
        self.assertNotIn("case_id", serialized_cells)
        self.assertNotIn("persona_id", serialized_cells)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
